from io import BytesIO
from decimal import Decimal
from abc import ABC, abstractmethod
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class AccountingExportAdapter(ABC):
    """Interfaz abstracta para exportación contable hacia sistemas ERP."""

    @abstractmethod
    def export(self, data: dict, from_date: str, to_date: str) -> bytes:
        pass


class ExcelAccountingExporter(AccountingExportAdapter):
    """Genera archivo Excel (.xlsx) oficial con desglose contable para Contabilidad."""

    def export(self, data: dict, from_date: str, to_date: str) -> bytes:
        wb = openpyxl.Workbook()

        # Estilos corporativos POTOQUITOS
        header_fill = PatternFill(start_color="1255A0", end_color="1255A0", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=15, bold=True, color="1255A0")
        subtitle_font = Font(name="Calibri", size=11, italic=True, color="4A5568")
        bold_font = Font(name="Calibri", size=11, bold=True)
        total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E0"),
            right=Side(style="thin", color="CBD5E0"),
            top=Side(style="thin", color="CBD5E0"),
            bottom=Side(style="thin", color="CBD5E0"),
        )
        total_border = Border(
            top=Side(style="thin", color="000000"),
            bottom=Side(style="double", color="000000"),
        )

        # -------------------------------------------------------------
        # HOJA 1: VENTAS Y COBROS
        # -------------------------------------------------------------
        ws_ventas = wb.active
        ws_ventas.title = "Ventas y Cobros"
        ws_ventas.views.sheetView[0].showGridLines = True

        ws_ventas["A1"] = "POTOQUITOS — EXPORTACIÓN CONTABLE DE VENTAS"
        ws_ventas["A1"].font = title_font
        ws_ventas["A2"] = f"Período: {from_date} al {to_date}"
        ws_ventas["A2"].font = subtitle_font

        headers_ventas = [
            "Fecha",
            "Hora",
            "Factura / Comp.",
            "Mesa",
            "Cajero",
            "Turno / Caja",
            "Concepto / Ítems",
            "Consumo / Venta (COP)",
            "Propina (COP)",
            "Impuestos (COP)",
            "Total Cobrado (COP)",
            "Medio de Pago",
            "Efectivo (COP)",
            "Tarjeta (COP)",
            "Transferencia (COP)",
        ]

        for col_num, h_text in enumerate(headers_ventas, 1):
            cell = ws_ventas.cell(row=4, column=col_num, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        row_idx = 5
        tot_consumo = Decimal("0.00")
        tot_propina = Decimal("0.00")
        tot_impuesto = Decimal("0.00")
        tot_cobrado = Decimal("0.00")
        tot_efectivo = Decimal("0.00")
        tot_tarjeta = Decimal("0.00")
        tot_transfer = Decimal("0.00")

        payments = data.get("payments", [])
        for p in payments:
            cons = Decimal(str(p.get("consumption_amount", 0)))
            prop = Decimal(str(p.get("tip_amount", 0)))
            imp = Decimal(str(p.get("tax_amount", 0)))
            tot = Decimal(str(p.get("total_amount", 0)))
            efec = Decimal(str(p.get("cash_amount", 0)))
            tarj = Decimal(str(p.get("card_amount", 0)))
            trans = Decimal(str(p.get("transfer_amount", 0)))

            tot_consumo += cons
            tot_propina += prop
            tot_impuesto += imp
            tot_cobrado += tot
            tot_efectivo += efec
            tot_tarjeta += tarj
            tot_transfer += trans

            values = [
                p.get("date", ""),
                p.get("time", ""),
                p.get("invoice_number", ""),
                p.get("table_number", ""),
                p.get("cashier_name", ""),
                p.get("register_name", ""),
                p.get("concept", "Consumo en salón"),
                float(cons),
                float(prop),
                float(imp),
                float(tot),
                p.get("payment_method_summary", ""),
                float(efec),
                float(tarj),
                float(trans),
            ]

            for c_idx, val in enumerate(values, 1):
                cell = ws_ventas.cell(row=row_idx, column=c_idx, value=val)
                cell.border = thin_border
                if c_idx in (8, 9, 10, 11, 13, 14, 15):
                    cell.number_format = "$#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif c_idx in (1, 2, 3, 4):
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(horizontal="left")
            row_idx += 1

        # Fila de Totales Ventas
        ws_ventas.cell(row=row_idx, column=1, value="TOTALES DEL PERÍODO").font = bold_font
        for c in range(1, 16):
            ws_ventas.cell(row=row_idx, column=c).fill = total_fill
            ws_ventas.cell(row=row_idx, column=c).border = total_border

        ws_ventas.cell(row=row_idx, column=8, value=float(tot_consumo)).number_format = "$#,##0.00"
        ws_ventas.cell(row=row_idx, column=9, value=float(tot_propina)).number_format = "$#,##0.00"
        ws_ventas.cell(row=row_idx, column=10, value=float(tot_impuesto)).number_format = "$#,##0.00"
        ws_ventas.cell(row=row_idx, column=11, value=float(tot_cobrado)).number_format = "$#,##0.00"
        ws_ventas.cell(row=row_idx, column=13, value=float(tot_efectivo)).number_format = "$#,##0.00"
        ws_ventas.cell(row=row_idx, column=14, value=float(tot_tarjeta)).number_format = "$#,##0.00"
        ws_ventas.cell(row=row_idx, column=15, value=float(tot_transfer)).number_format = "$#,##0.00"

        # -------------------------------------------------------------
        # HOJA 2: GASTOS Y EGRESOS
        # -------------------------------------------------------------
        ws_gastos = wb.create_sheet(title="Gastos y Egresos")
        ws_gastos.views.sheetView[0].showGridLines = True

        ws_gastos["A1"] = "POTOQUITOS — REGISTRO DE GASTOS Y COSTOS"
        ws_gastos["A1"].font = title_font
        ws_gastos["A2"] = f"Período: {from_date} al {to_date}"
        ws_gastos["A2"].font = subtitle_font

        headers_gastos = [
            "Fecha",
            "Hora",
            "Categoría de Gasto",
            "Descripción / Justificación",
            "Proveedor / Referencia",
            "Valor (COP)",
            "Medio de Pago",
            "Registrado Por",
        ]

        for col_num, h_text in enumerate(headers_gastos, 1):
            cell = ws_gastos.cell(row=4, column=col_num, value=h_text)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        g_row_idx = 5
        tot_gastos = Decimal("0.00")
        expenses = data.get("expenses", [])
        for e in expenses:
            val = Decimal(str(e.get("amount", 0)))
            tot_gastos += val
            g_values = [
                e.get("date", ""),
                e.get("time", ""),
                e.get("category_name", ""),
                e.get("description", ""),
                e.get("reference", ""),
                float(val),
                e.get("payment_method", "EFECTIVO"),
                e.get("responsible_name", ""),
            ]
            for c_idx, val_item in enumerate(g_values, 1):
                cell = ws_gastos.cell(row=g_row_idx, column=c_idx, value=val_item)
                cell.border = thin_border
                if c_idx == 6:
                    cell.number_format = "$#,##0.00"
                    cell.alignment = Alignment(horizontal="right")
                elif c_idx in (1, 2):
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(horizontal="left")
            g_row_idx += 1

        ws_gastos.cell(row=g_row_idx, column=1, value="TOTAL GASTOS").font = bold_font
        for c in range(1, 9):
            ws_gastos.cell(row=g_row_idx, column=c).fill = total_fill
            ws_gastos.cell(row=g_row_idx, column=c).border = total_border
        ws_gastos.cell(row=g_row_idx, column=6, value=float(tot_gastos)).number_format = "$#,##0.00"

        # -------------------------------------------------------------
        # HOJA 3: RESUMEN CONTABLE
        # -------------------------------------------------------------
        ws_resumen = wb.create_sheet(title="Resumen Contable")
        ws_resumen.views.sheetView[0].showGridLines = True

        ws_resumen["A1"] = "POTOQUITOS — RESUMEN FINANCIERO Y CONTABLE"
        ws_resumen["A1"].font = title_font
        ws_resumen["A2"] = f"Período: {from_date} al {to_date}"
        ws_resumen["A2"].font = subtitle_font

        ws_resumen.cell(row=4, column=1, value="Concepto Contable").font = header_font
        ws_resumen.cell(row=4, column=1).fill = header_fill
        ws_resumen.cell(row=4, column=2, value="Valor Total (COP)").font = header_font
        ws_resumen.cell(row=4, column=2).fill = header_fill

        summary_rows = [
            ("Ventas / Consumo de Alimentos y Bebidas", float(tot_consumo)),
            ("Propinas Voluntarias Recibidas", float(tot_propina)),
            ("Impuestos Generados", float(tot_impuesto)),
            ("TOTAL FACTURACIÓN BRUTA", float(tot_cobrado)),
            ("  - Recaudado en Efectivo", float(tot_efectivo)),
            ("  - Recaudado en Tarjeta Débito/Crédito", float(tot_tarjeta)),
            ("  - Recaudado en Transferencia Bancaria", float(tot_transfer)),
            ("TOTAL GASTOS Y EGRESOS", float(tot_gastos)),
            ("FLUJO NETO OPERATIVO DEL PERÍODO", float(tot_cobrado - tot_gastos)),
        ]

        r_idx = 5
        for label, amount in summary_rows:
            c1 = ws_resumen.cell(row=r_idx, column=1, value=label)
            c2 = ws_resumen.cell(row=r_idx, column=2, value=amount)
            c1.border = thin_border
            c2.border = thin_border
            c2.number_format = "$#,##0.00"
            c2.alignment = Alignment(horizontal="right")
            if "TOTAL" in label or "FLUJO" in label:
                c1.font = bold_font
                c2.font = bold_font
                c1.fill = total_fill
                c2.fill = total_fill
            r_idx += 1

        # Ajuste automático de anchos de columna para todas las hojas
        for ws in (ws_ventas, ws_gastos, ws_resumen):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        out = BytesIO()
        wb.save(out)
        return out.getvalue()


class XmlAccountingExporter(AccountingExportAdapter):
    """Adaptador extensible preparado para exportaciones XML hacia software contables (Siigo, Helisa, etc.)."""

    def export(self, data: dict, from_date: str, to_date: str) -> bytes:
        payments = data.get("payments", [])
        expenses = data.get("expenses", [])

        xml_lines = [
            '<?xml version="1.0" encoding="UTF-8"?>',
            f'<PotoquitosAccountingExport period_from="{from_date}" period_to="{to_date}">',
            '  <Establishment name="POTOQUITOS" nit="901.458.789-2" />',
            '  <Invoices>',
        ]

        for p in payments:
            xml_lines.append(
                f'    <Invoice number="{p.get("invoice_number")}" date="{p.get("date")}" time="{p.get("time")}" table="{p.get("table_number")}">'
            )
            xml_lines.append(f'      <Consumption amount="{p.get("consumption_amount", 0)}" />')
            xml_lines.append(f'      <Tip amount="{p.get("tip_amount", 0)}" />')
            xml_lines.append(f'      <Total amount="{p.get("total_amount", 0)}" />')
            xml_lines.append(f'      <PaymentMethod method="{p.get("payment_method_summary", "")}" />')
            xml_lines.append('    </Invoice>')

        xml_lines.append('  </Invoices>')
        xml_lines.append('  <Expenses>')
        for e in expenses:
            xml_lines.append(
                f'    <Expense date="{e.get("date")}" category="{e.get("category_name")}" amount="{e.get("amount")}" method="{e.get("payment_method")}" />'
            )
        xml_lines.append('  </Expenses>')
        xml_lines.append('</PotoquitosAccountingExport>')

        return "\n".join(xml_lines).encode("utf-8")
